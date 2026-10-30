from __future__ import annotations

import csv
import io
import os
import re
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from typing import Any

import socketio
from fastapi import Depends, FastAPI, HTTPException, Query, Request, Response, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import PlainTextResponse, StreamingResponse
from prometheus_client import CONTENT_TYPE_LATEST, Counter, Histogram, generate_latest
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import and_, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload
from starlette.concurrency import run_in_threadpool
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer

from .database import build_engine, build_session_factory
from .models import (
    Base, Bitacora, Categoria, DetalleOrden, Insumo, Inventario, Mesa,
    MovimientoInventario, Orden, Pago, Permiso, Producto, RefreshToken,
    Rol, Sucursal, Usuario, recetas,
)
from .security import (
    ACCESS_MINUTES, REFRESH_DAYS, create_access_token, decode_access_token,
    hash_password, hash_refresh_token, new_refresh_token, refresh_expiry,
    verify_password,
)

API_REQUESTS = Counter("comandapro_http_requests_total", "HTTP requests handled", ["method", "path", "status"])
API_LATENCY = Histogram("comandapro_http_request_duration_seconds", "HTTP request latency", ["method", "path"])
SOCKET_ALLOWED_ORIGINS = [x.strip() for x in os.getenv(
    "CORS_ORIGINS",
    "http://localhost:8080,http://127.0.0.1:8080,http://localhost:5173",
).split(",") if x.strip()]
RENDER_ORIGIN = os.getenv("RENDER_EXTERNAL_URL", "").strip()
if RENDER_ORIGIN and RENDER_ORIGIN not in SOCKET_ALLOWED_ORIGINS:
    SOCKET_ALLOWED_ORIGINS.append(RENDER_ORIGIN)
SOCKET = socketio.AsyncServer(async_mode="asgi", cors_allowed_origins=SOCKET_ALLOWED_ORIGINS, logger=False, engineio_logger=False)
_SOCKET_SESSION_FACTORY = None

ALL_PERMISSIONS: dict[str, str] = {
    "sucursales.ver": "Consultar sucursales",
    "sucursales.ver_todas": "Asignar usuarios y consultar sucursales de la cadena",
    "sucursales.crear": "Crear y editar sucursales",
    "usuarios.ver": "Consultar usuarios de la sucursal",
    "usuarios.ver_todas": "Consultar y administrar usuarios de toda la cadena",
    "usuarios.crear": "Crear usuarios",
    "usuarios.editar": "Editar o desactivar usuarios",
    "roles.configurar": "Configurar roles y permisos",
    "productos.ver": "Consultar menú y productos",
    "productos.crear": "Crear productos y categorías",
    "productos.editar": "Editar productos y categorías",
    "productos.eliminar": "Desactivar productos",
    "recetas.configurar": "Configurar recetas de productos",
    "insumos.ver": "Consultar insumos e inventario",
    "insumos.crear": "Crear insumos",
    "inventario.ver": "Consultar existencias de su sucursal",
    "inventario.ajustar": "Registrar entradas y ajustes de inventario",
    "mesas.ver": "Consultar mesas de su sucursal",
    "mesas.configurar": "Crear y editar mesas",
    "ordenes.ver": "Consultar órdenes de su sucursal",
    "ordenes.crear": "Crear órdenes y detalles",
    "ordenes.editar": "Editar órdenes abiertas",
    "ordenes.enviar_cocina": "Enviar órdenes a cocina",
    "ordenes.cancelar": "Cancelar órdenes o platillos",
    "cocina.ver": "Ver comandas en cocina",
    "cocina.actualizar": "Actualizar estado de platillos",
    "pagos.ver": "Consultar cuentas para cobro",
    "pagos.crear": "Registrar pagos y pagos divididos",
    "reportes.ver": "Consultar reportes de su sucursal",
    "reportes.ver_todas": "Comparar ventas entre sucursales",
    "bitacora.ver": "Consultar la bitácora de auditoría",
}
ROLE_PERMISSIONS = {
    "ADMIN": list(ALL_PERMISSIONS),
    "MESERO": ["mesas.ver", "productos.ver", "ordenes.ver", "ordenes.crear", "ordenes.editar", "ordenes.enviar_cocina", "ordenes.cancelar"],
    "COCINA": ["productos.ver", "ordenes.ver", "cocina.ver", "cocina.actualizar"],
    "CAJERO": ["ordenes.ver", "pagos.ver", "pagos.crear", "reportes.ver"],
}


def money(value: Any) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _payload(obj: Any) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in obj.__dict__.items():
        if key.startswith("_"):
            continue
        if isinstance(value, Decimal):
            result[key] = float(value)
        elif isinstance(value, datetime):
            result[key] = value.isoformat()
        elif hasattr(value, "value"):
            result[key] = value.value
        else:
            result[key] = value
    return result


def _user_payload(user: Usuario, permissions: list[str] | None = None) -> dict[str, Any]:
    return {
        "id": user.id, "nombre": user.nombre, "correo": user.correo,
        "rol": user.rol.nombre if user.rol else "", "rol_id": user.rol_id,
        "sucursal_id": user.sucursal_id,
        "sucursal": user.sucursal.nombre if user.sucursal else "",
        "activo": user.activo, "permisos": permissions or [],
    }


class LoginInput(BaseModel):
    correo: str
    password: str


class PasswordChangeInput(BaseModel):
    actual: str
    nueva: str = Field(min_length=16, max_length=128)


class RoleInput(BaseModel):
    nombre: str = Field(min_length=2, max_length=40, pattern="^[A-Za-zÁÉÍÓÚÑáéíóúñ_ ]+$")


class PermissionInput(BaseModel):
    clave: str = Field(min_length=3, max_length=100, pattern="^[a-z]+([._][a-z0-9_]+)+$")
    descripcion: str = Field(min_length=3, max_length=240)


class BranchUpdate(BaseModel):
    nombre: str | None = Field(default=None, min_length=2, max_length=120)
    direccion: str | None = Field(default=None, max_length=240)
    activa: bool | None = None


class UserCreate(BaseModel):
    nombre: str = Field(min_length=2, max_length=140)
    correo: str = Field(min_length=5, max_length=180)
    password: str = Field(min_length=10, max_length=128)
    rol_id: int
    sucursal_id: int | None = None


class UserUpdate(BaseModel):
    nombre: str | None = Field(default=None, min_length=2, max_length=140)
    correo: str | None = Field(default=None, min_length=5, max_length=180)
    rol_id: int | None = None
    activo: bool | None = None


class SucursalInput(BaseModel):
    nombre: str = Field(min_length=2, max_length=120)
    direccion: str = Field(default="", max_length=240)
    activa: bool = True


class MesaInput(BaseModel):
    numero: int = Field(ge=1, le=1000)
    capacidad: int = Field(default=4, ge=1, le=30)
    estado: str = "LIBRE"


class CategoriaInput(BaseModel):
    nombre: str = Field(min_length=2, max_length=100)


class ProductoInput(BaseModel):
    categoria_id: int
    nombre: str = Field(min_length=2, max_length=140)
    precio: Decimal = Field(gt=0, max_digits=10, decimal_places=2)
    imagen_url: str = Field(default="", max_length=500)
    activo: bool = True


class RecetaInput(BaseModel):
    insumo_id: int
    cantidad: Decimal = Field(gt=0, max_digits=12, decimal_places=3)


class RecetaReplace(BaseModel):
    insumos: list[RecetaInput]


class InsumoInput(BaseModel):
    nombre: str = Field(min_length=2, max_length=140)
    unidad_medida: str = Field(min_length=1, max_length=24)


class InventarioAdjustment(BaseModel):
    cantidad: Decimal = Field(gt=0, max_digits=12, decimal_places=3)
    tipo: str = Field(pattern="^(ENTRADA|AJUSTE_POSITIVO|AJUSTE_NEGATIVO)$")
    nota: str = Field(default="", max_length=300)


class OrdenLineInput(BaseModel):
    producto_id: int
    cantidad: int = Field(ge=1, le=100)
    notas: str = Field(default="", max_length=300)


class OrdenCreate(BaseModel):
    mesa_id: int
    detalles: list[OrdenLineInput] = Field(min_length=1, max_length=40)


class OrdenLineAdd(BaseModel):
    producto_id: int
    cantidad: int = Field(ge=1, le=100)
    notas: str = Field(default="", max_length=300)


class OrdenLineUpdate(BaseModel):
    cantidad: int | None = Field(default=None, ge=1, le=100)
    notas: str | None = Field(default=None, max_length=300)


class DishStateInput(BaseModel):
    estado: str = Field(pattern="^(PREPARANDO|LISTO)$")


class PaymentInput(BaseModel):
    metodo: str = Field(pattern="^(EFECTIVO|TARJETA|TRANSFERENCIA)$")
    monto: Decimal = Field(gt=0, max_digits=10, decimal_places=2)


class RolePermissionsInput(BaseModel):
    permiso_ids: list[int]


class OrderCancelInput(BaseModel):
    motivo: str = Field(default="Cancelación solicitada", max_length=300)


def _seed_database(session: Session, admin_email: str, admin_password: str) -> None:
    if not session.query(Sucursal).count():
        session.add_all([
            Sucursal(nombre="ComandaPro Centro", direccion="Av. Principal 100"),
            Sucursal(nombre="ComandaPro Norte", direccion="Blvd. Norte 250"),
            Sucursal(nombre="ComandaPro Sur", direccion="Calle del Parque 58"),
        ])
        session.flush()
    roles: dict[str, Rol] = {}
    new_roles: set[str] = set()
    for role_name in ROLE_PERMISSIONS:
        role = session.query(Rol).filter_by(nombre=role_name).one_or_none()
        if role is None:
            role = Rol(nombre=role_name)
            session.add(role)
            session.flush()
            new_roles.add(role_name)
        roles[role_name] = role
    permissions: dict[str, Permiso] = {}
    for key, description in ALL_PERMISSIONS.items():
        permission = session.query(Permiso).filter_by(clave=key).one_or_none()
        if permission is None:
            permission = Permiso(clave=key, descripcion=description)
            session.add(permission)
            session.flush()
        permissions[key] = permission
    for role_name, keys in ROLE_PERMISSIONS.items():
        # Seed default grants only for new roles; keep administrator-configured permissions on later restarts.
        if role_name in new_roles:
            roles[role_name].permisos = [permissions[key] for key in keys]
    session.flush()
    if not session.query(Categoria).count():
        cats = [Categoria(nombre="Entradas"), Categoria(nombre="Platos fuertes"), Categoria(nombre="Bebidas"), Categoria(nombre="Postres")]
        session.add_all(cats)
        session.flush()
    if not session.query(Insumo).count():
        session.add_all([
            Insumo(nombre="Pan artesanal", unidad_medida="pieza"),
            Insumo(nombre="Carne de res", unidad_medida="kg"),
            Insumo(nombre="Queso", unidad_medida="kg"),
            Insumo(nombre="Jitomate", unidad_medida="kg"),
            Insumo(nombre="Lechuga", unidad_medida="kg"),
            Insumo(nombre="Papa", unidad_medida="kg"),
            Insumo(nombre="Refresco", unidad_medida="pieza"),
            Insumo(nombre="Harina", unidad_medida="kg"),
        ])
        session.flush()
    if not session.query(Producto).count():
        c = {x.nombre: x.id for x in session.query(Categoria).all()}
        session.add_all([
            Producto(categoria_id=c["Platos fuertes"], nombre="Hamburguesa Comanda", precio=money("125.00"), imagen_url="", activo=True),
            Producto(categoria_id=c["Platos fuertes"], nombre="Papas de la casa", precio=money("55.00"), imagen_url="", activo=True),
            Producto(categoria_id=c["Bebidas"], nombre="Refresco", precio=money("32.00"), imagen_url="", activo=True),
            Producto(categoria_id=c["Entradas"], nombre="Pan con queso", precio=money("48.00"), imagen_url="", activo=True),
        ])
        session.flush()
    if not session.execute(select(func.count()).select_from(recetas)).scalar_one():
        prod = {x.nombre: x.id for x in session.query(Producto).all()}
        ins = {x.nombre: x.id for x in session.query(Insumo).all()}
        recipes = [
            {"producto_id": prod["Hamburguesa Comanda"], "insumo_id": ins["Pan artesanal"], "cantidad": Decimal("1")},
            {"producto_id": prod["Hamburguesa Comanda"], "insumo_id": ins["Carne de res"], "cantidad": Decimal("0.180")},
            {"producto_id": prod["Hamburguesa Comanda"], "insumo_id": ins["Queso"], "cantidad": Decimal("0.030")},
            {"producto_id": prod["Hamburguesa Comanda"], "insumo_id": ins["Jitomate"], "cantidad": Decimal("0.040")},
            {"producto_id": prod["Hamburguesa Comanda"], "insumo_id": ins["Lechuga"], "cantidad": Decimal("0.020")},
            {"producto_id": prod["Papas de la casa"], "insumo_id": ins["Papa"], "cantidad": Decimal("0.250")},
            {"producto_id": prod["Refresco"], "insumo_id": ins["Refresco"], "cantidad": Decimal("1")},
            {"producto_id": prod["Pan con queso"], "insumo_id": ins["Pan artesanal"], "cantidad": Decimal("1")},
            {"producto_id": prod["Pan con queso"], "insumo_id": ins["Queso"], "cantidad": Decimal("0.050")},
        ]
        session.execute(recetas.insert(), recipes)
    if not session.query(Inventario).count():
        ingredients = session.query(Insumo).all()
        for branch in session.query(Sucursal).all():
            for ing in ingredients:
                q = Decimal("40") if ing.unidad_medida == "pieza" else Decimal("8.000")
                session.add(Inventario(sucursal_id=branch.id, insumo_id=ing.id, existencia=q, stock_minimo=Decimal("2")))
    if not session.query(Mesa).count():
        for branch in session.query(Sucursal).all():
            for n in range(1, 13):
                session.add(Mesa(sucursal_id=branch.id, numero=n, capacidad=4 if n <= 8 else 6, estado="LIBRE"))
    session.flush()
    admin_user = session.query(Usuario).filter_by(correo=admin_email.lower()).one_or_none()
    if admin_user is None:
        session.add(Usuario(nombre="Administración ComandaPro", correo=admin_email.lower(), password_hash=hash_password(admin_password), rol_id=roles["ADMIN"].id, sucursal_id=1, activo=True))
    elif (
        os.getenv("COMANDAPRO_ADMIN_PASSWORD")
        and admin_password != "CambiaComanda2026!"
        and verify_password("CambiaComanda2026!", admin_user.password_hash)
    ):
        # Migra únicamente la contraseña inicial predeterminada.
        # No sustituye una contraseña que ya haya sido personalizada.
        admin_user.password_hash = hash_password(admin_password)
    demo_password = os.getenv("COMANDAPRO_DEMO_PASSWORD", "DemoComanda2026!")
    for role_name, email_prefix in (("MESERO", "mesero"), ("COCINA", "cocina"), ("CAJERO", "cajero")):
        email = f"{email_prefix}@comandapro.local"
        if not session.query(Usuario).filter_by(correo=email).one_or_none():
            session.add(Usuario(nombre=f"Usuario {role_name.title()}", correo=email, password_hash=hash_password(demo_password), rol_id=roles[role_name].id, sucursal_id=1, activo=True))
    session.commit()


AUDIT_NAMES = {"productos": "producto", "categorias": "categoria", "insumos": "insumo", "mesas": "mesa", "usuarios": "usuario", "roles": "rol", "permisos": "permiso", "sucursales": "sucursal"}
AUDIT_VERBS = {"CREAR": "Creó", "EDITAR": "Editó", "ELIMINAR": "Eliminó"}


def describe_audit(method: str, path: str) -> tuple[str, str, int | None, str]:
    parts = [p for p in path.strip("/").split("/") if p and p != "api"]
    ids = [int(p) for p in parts if p.isdigit()]
    names = [p for p in parts if not p.isdigit()]
    ref = ids[0] if ids else None
    action = {"POST": "CREAR", "PUT": "EDITAR", "PATCH": "EDITAR", "DELETE": "ELIMINAR"}.get(method, method)
    if names == ["auth", "login"]:
        return "ACCESO", "sesion", None, "Inició sesión"
    if names[:1] == ["ordenes"]:
        tail = names[1:]
        if not tail:
            return "CREAR", "orden", None, "Creó una orden"
        if tail == ["enviar-cocina"]:
            return "EDITAR", "orden", ref, f"Envió la orden #{ref} a cocina"
        if tail == ["servir"]:
            return "EDITAR", "orden", ref, f"Marcó la orden #{ref} como servida"
        if tail == ["cancelar"]:
            return "EDITAR", "orden", ref, f"Canceló la orden #{ref}"
        if tail == ["pagos"]:
            return "CREAR", "pago", ref, f"Registró un pago de la orden #{ref}"
        if tail == ["detalles"]:
            if len(ids) == 1:
                return "CREAR", "platillo", ref, f"Agregó un platillo a la orden #{ref}"
            detail_id = ids[1]
            if method == "DELETE":
                return "ELIMINAR", "platillo", detail_id, f"Canceló el platillo #{detail_id} de la orden #{ref}"
            return "EDITAR", "platillo", detail_id, f"Modificó el platillo #{detail_id} de la orden #{ref}"
    if names[:1] == ["cocina"]:
        return "EDITAR", "platillo", ref, f"Cambió el estado del platillo #{ref} en cocina"
    if names[:1] == ["inventario"]:
        return "CREAR", "movimiento_inventario", ref, f"Registró un movimiento del inventario #{ref}"
    if names[-1:] == ["receta"]:
        return "EDITAR", "receta", ref, f"Actualizó la receta del producto #{ref}"
    if names[-1:] == ["permisos"] and ids:
        return "EDITAR", "rol", ref, f"Actualizó los permisos del rol #{ref}"
    noun = AUDIT_NAMES.get(names[-1], names[-1]) if names else "registro"
    suffix = f" #{ref}" if ref else ""
    return action, noun, ref, f"{AUDIT_VERBS.get(action, action)} {noun}{suffix}"


def create_app(database_url: str | None = None, initialize: bool = True) -> FastAPI:
    engine = build_engine(database_url)
    SessionLocal = build_session_factory(engine)
    app = FastAPI(title="ComandaPro API", version="1.0.0", description="API de pedidos multi-sucursal para restaurantes", docs_url="/api/docs", openapi_url="/api/openapi.json")
    app.state.engine = engine
    app.state.SessionLocal = SessionLocal
    app.state.initialize_db = initialize
    app.add_middleware(CORSMiddleware, allow_origins=SOCKET_ALLOWED_ORIGINS, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

    @app.on_event("startup")
    def startup() -> None:
        if not app.state.initialize_db:
            return
        if os.getenv("COMANDAPRO_SERVE_FRONTEND", "").lower() == "true":
            admin_secret = os.getenv("COMANDAPRO_ADMIN_PASSWORD", "")
            if len(admin_secret) < 16 or admin_secret == "CambiaComanda2026!":
                raise RuntimeError("El despliegue público requiere COMANDAPRO_ADMIN_PASSWORD seguro de al menos 16 caracteres")
        Base.metadata.create_all(bind=engine)
        with SessionLocal() as db:
            _seed_database(db, os.getenv("COMANDAPRO_ADMIN_EMAIL", "admin@comandapro.local"), os.getenv("COMANDAPRO_ADMIN_PASSWORD", "CambiaComanda2026!"))

    def get_db(request: Request):
        db = request.app.state.SessionLocal()
        try:
            yield db
        finally:
            db.close()

    bearer = HTTPBearer(auto_error=False)

    def permissions_for(db: Session, user: Usuario) -> list[str]:
        return sorted({p.clave for p in user.rol.permisos}) if user.rol else []

    def current_user(request: Request, creds: HTTPAuthorizationCredentials | None = Depends(bearer), db: Session = Depends(get_db)) -> Usuario:
        if not creds:
            raise HTTPException(status_code=401, detail="Se requiere un token Bearer")
        try:
            payload = decode_access_token(creds.credentials)
            uid = int(payload["sub"])
        except Exception:
            raise HTTPException(status_code=401, detail="Token de acceso inválido o expirado")
        user = db.query(Usuario).options(joinedload(Usuario.rol), joinedload(Usuario.sucursal)).filter(Usuario.id == uid, Usuario.activo.is_(True)).one_or_none()
        if user is None:
            raise HTTPException(status_code=401, detail="Usuario inactivo o inexistente")
        request.state.user_id = user.id
        request.state.user_branch_id = user.sucursal_id
        request.state.user_permissions = permissions_for(db, user)
        return user

    def require_permission(key: str):
        def dependency(user: Usuario = Depends(current_user), db: Session = Depends(get_db)) -> Usuario:
            if key not in permissions_for(db, user):
                raise HTTPException(status_code=403, detail=f"Falta el permiso requerido: {key}")
            return user
        return dependency

    def scoped_branch(user: Usuario, requested: int | None = None, allow_global: bool = False) -> int:
        perms = set([p.clave for p in user.rol.permisos]) if user.rol else set()
        if allow_global and "sucursales.ver_todas" in perms:
            return requested if requested is not None else user.sucursal_id
        if requested is not None and requested != user.sucursal_id:
            raise HTTPException(status_code=403, detail="No puede consultar otra sucursal")
        return user.sucursal_id

    def _branch_orders(db: Session, user: Usuario):
        return db.query(Orden).filter(Orden.sucursal_id == user.sucursal_id)

    @app.middleware("http")
    async def metrics_and_audit(request: Request, call_next):
        started = __import__("time").perf_counter()
        response = await call_next(request)
        elapsed = __import__("time").perf_counter() - started
        path = request.url.path
        # Normalize numeric resource IDs in metric labels to avoid time-series cardinality growth.
        metric_path = re.sub(r"/\d+(?=/|$)", "/{id}", path)
        API_REQUESTS.labels(request.method, metric_path, str(response.status_code)).inc()
        API_LATENCY.labels(request.method, metric_path).observe(elapsed)
        if request.method in {"POST", "PUT", "PATCH", "DELETE"} and response.status_code < 400:
            user_id = getattr(request.state, "user_id", None)
            if user_id and not path.endswith(("/auth/refresh", "/auth/logout")):
                try:
                    accion, entidad, entidad_id, descripcion = describe_audit(request.method, path)
                    entidad_id = getattr(request.state, "audit_ref", None) or entidad_id
                    with request.app.state.SessionLocal() as audit_db:
                        audit_db.add(Bitacora(usuario_id=user_id, accion=accion, entidad=entidad, entidad_id=entidad_id, datos_json={"descripcion": descripcion, "ruta": path, "metodo": request.method, "status": response.status_code}))
                        audit_db.commit()
                except Exception:
                    # An audit failure must not transform an otherwise successful API response.
                    pass
        return response

    @app.get("/api/health")
    def health(db: Session = Depends(get_db)):
        db.execute(select(1))
        return {"estado": "ok", "baseDatos": "conectada"}

    @app.get("/metrics", include_in_schema=False)
    def metrics():
        return Response(generate_latest(), media_type=CONTENT_TYPE_LATEST)

    @app.post("/api/auth/login")
    def login(body: LoginInput, response: Response, request: Request, db: Session = Depends(get_db)):
        email = body.correo.strip().lower()
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            raise HTTPException(status_code=422, detail="Correo inválido")
        user = db.query(Usuario).options(joinedload(Usuario.rol), joinedload(Usuario.sucursal)).filter(func.lower(Usuario.correo) == email).one_or_none()
        if not user or not user.activo or not verify_password(body.password, user.password_hash):
            raise HTTPException(status_code=401, detail="Correo o contraseña incorrectos")
        raw_refresh = new_refresh_token()
        db.add(RefreshToken(usuario_id=user.id, token_hash=hash_refresh_token(raw_refresh), expira_en=refresh_expiry(), revocado=False))
        db.commit()
        secure = os.getenv("COOKIE_SECURE", "false").lower() == "true"
        response.set_cookie("comandapro_refresh", raw_refresh, httponly=True, secure=secure, samesite="lax", max_age=REFRESH_DAYS * 86400, path="/api/auth")
        permissions = permissions_for(db, user)
        request.state.user_id = user.id
        request.state.user_branch_id = user.sucursal_id
        return {"access_token": create_access_token(user.id, user.sucursal_id), "token_type": "bearer", "expires_in": ACCESS_MINUTES * 60, "usuario": _user_payload(user, permissions)}

    @app.post("/api/auth/refresh")
    def refresh(request: Request, response: Response, db: Session = Depends(get_db)):
        token = request.cookies.get("comandapro_refresh")
        if not token:
            raise HTTPException(status_code=401, detail="No existe refresh token")
        record = db.query(RefreshToken).filter(RefreshToken.token_hash == hash_refresh_token(token)).with_for_update().one_or_none()
        now = datetime.now(timezone.utc)
        if not record or record.revocado or (record.expira_en.replace(tzinfo=timezone.utc) if record.expira_en.tzinfo is None else record.expira_en) <= now:
            raise HTTPException(status_code=401, detail="Refresh token revocado o expirado")
        user = db.query(Usuario).options(joinedload(Usuario.rol), joinedload(Usuario.sucursal)).filter(Usuario.id == record.usuario_id, Usuario.activo.is_(True)).one_or_none()
        if not user:
            record.revocado = True
            db.commit()
            raise HTTPException(status_code=401, detail="Usuario inactivo")
        new_raw = new_refresh_token()
        new_hash = hash_refresh_token(new_raw)
        record.revocado = True
        record.reemplaza_a = new_hash
        db.add(RefreshToken(usuario_id=user.id, token_hash=new_hash, expira_en=refresh_expiry(), revocado=False))
        db.commit()
        secure = os.getenv("COOKIE_SECURE", "false").lower() == "true"
        response.set_cookie("comandapro_refresh", new_raw, httponly=True, secure=secure, samesite="lax", max_age=REFRESH_DAYS * 86400, path="/api/auth")
        return {"access_token": create_access_token(user.id, user.sucursal_id), "token_type": "bearer", "expires_in": ACCESS_MINUTES * 60, "usuario": _user_payload(user, permissions_for(db, user))}

    @app.post("/api/auth/logout", status_code=204)
    def logout(request: Request, response: Response, db: Session = Depends(get_db), user: Usuario = Depends(current_user)):
        token = request.cookies.get("comandapro_refresh")
        if token:
            record = db.query(RefreshToken).filter(RefreshToken.token_hash == hash_refresh_token(token), RefreshToken.usuario_id == user.id).one_or_none()
            if record:
                record.revocado = True
                db.commit()
        response.delete_cookie("comandapro_refresh", path="/api/auth")
        response.status_code = 204
        response.body = b""
        return response

    @app.post("/api/auth/change-password")
    def change_password(body: PasswordChangeInput, response: Response, user: Usuario = Depends(current_user), db: Session = Depends(get_db)):
        if not verify_password(body.actual, user.password_hash):
            raise HTTPException(status_code=401, detail="La contraseña actual es incorrecta")
        if body.actual == body.nueva:
            raise HTTPException(status_code=400, detail="La nueva contraseña debe ser diferente")
        user.password_hash = hash_password(body.nueva)
        db.query(RefreshToken).filter(RefreshToken.usuario_id == user.id, RefreshToken.revocado.is_(False)).update(
            {RefreshToken.revocado: True}, synchronize_session=False
        )
        db.commit()
        response.delete_cookie("comandapro_refresh", path="/api/auth")
        return {"mensaje": "Contraseña actualizada. Inicia sesión nuevamente."}

    @app.get("/api/auth/me")
    def me(user: Usuario = Depends(current_user), db: Session = Depends(get_db)):
        return _user_payload(user, permissions_for(db, user))

    # Branches, users, roles and configurable permissions.
    @app.get("/api/admin/sucursales")
    def list_branches(user: Usuario = Depends(require_permission("sucursales.ver")), db: Session = Depends(get_db)):
        return [_payload(x) for x in db.query(Sucursal).order_by(Sucursal.id).all()]

    @app.post("/api/admin/sucursales", status_code=201)
    def create_branch(body: SucursalInput, user: Usuario = Depends(require_permission("sucursales.crear")), db: Session = Depends(get_db)):
        branch = Sucursal(**body.model_dump())
        try:
            db.add(branch)
            db.flush()
            # New branches receive their own table map and per-ingredient inventory rows.
            for number in range(1, 13):
                db.add(Mesa(sucursal_id=branch.id, numero=number, capacidad=4 if number <= 8 else 6, estado="LIBRE"))
            for ingredient in db.query(Insumo).all():
                db.add(Inventario(sucursal_id=branch.id, insumo_id=ingredient.id, existencia=Decimal("0"), stock_minimo=Decimal("2")))
            db.commit()
        except IntegrityError:
            db.rollback(); raise HTTPException(status_code=409, detail="Ya existe esa sucursal")
        db.refresh(branch)
        return _payload(branch)

    @app.patch("/api/admin/sucursales/{branch_id}")
    def update_branch(branch_id: int, body: BranchUpdate, user: Usuario = Depends(require_permission("sucursales.crear")), db: Session = Depends(get_db)):
        branch = db.get(Sucursal, branch_id)
        if not branch: raise HTTPException(status_code=404, detail="Sucursal no encontrada")
        for key, value in body.model_dump(exclude_unset=True).items(): setattr(branch, key, value)
        try: db.commit()
        except IntegrityError: db.rollback(); raise HTTPException(status_code=409, detail="Nombre de sucursal duplicado")
        db.refresh(branch); return _payload(branch)

    @app.get("/api/admin/roles")
    def list_roles(user: Usuario = Depends(require_permission("roles.configurar")), db: Session = Depends(get_db)):
        return [{"id": role.id, "nombre": role.nombre, "permisos": [{"id": p.id, "clave": p.clave, "descripcion": p.descripcion} for p in role.permisos]} for role in db.query(Rol).order_by(Rol.id).all()]

    @app.post("/api/admin/roles", status_code=201)
    def create_role(body: RoleInput, user: Usuario = Depends(require_permission("roles.configurar")), db: Session = Depends(get_db)):
        role = Rol(nombre=body.nombre.strip().upper()); db.add(role)
        try: db.commit()
        except IntegrityError: db.rollback(); raise HTTPException(status_code=409, detail="El rol ya existe")
        db.refresh(role); return {"id": role.id, "nombre": role.nombre, "permisos": []}

    @app.patch("/api/admin/roles/{role_id}")
    def rename_role(role_id: int, body: RoleInput, user: Usuario = Depends(require_permission("roles.configurar")), db: Session = Depends(get_db)):
        role = db.get(Rol, role_id)
        if not role: raise HTTPException(status_code=404, detail="Rol no encontrado")
        role.nombre = body.nombre.strip().upper()
        try: db.commit()
        except IntegrityError: db.rollback(); raise HTTPException(status_code=409, detail="El rol ya existe")
        db.refresh(role); return {"id": role.id, "nombre": role.nombre, "permisos": sorted(p.clave for p in role.permisos)}

    @app.delete("/api/admin/roles/{role_id}", status_code=204)
    def delete_role(role_id: int, user: Usuario = Depends(require_permission("roles.configurar")), db: Session = Depends(get_db)):
        role = db.get(Rol, role_id)
        if not role: raise HTTPException(status_code=404, detail="Rol no encontrado")
        if role.nombre == "ADMIN": raise HTTPException(status_code=409, detail="No se puede eliminar el rol ADMIN")
        if db.query(Usuario).filter_by(rol_id=role.id).count(): raise HTTPException(status_code=409, detail="El rol tiene usuarios asignados")
        db.delete(role); db.commit(); return Response(status_code=204)

    @app.get("/api/admin/permisos")
    def list_permissions(user: Usuario = Depends(require_permission("roles.configurar")), db: Session = Depends(get_db)):
        return [{"id": p.id, "clave": p.clave, "descripcion": p.descripcion} for p in db.query(Permiso).order_by(Permiso.clave).all()]

    @app.post("/api/admin/permisos", status_code=201)
    def create_permission(body: PermissionInput, user: Usuario = Depends(require_permission("roles.configurar")), db: Session = Depends(get_db)):
        obj = Permiso(clave=body.clave.strip(), descripcion=body.descripcion.strip()); db.add(obj)
        try: db.commit()
        except IntegrityError: db.rollback(); raise HTTPException(status_code=409, detail="La clave de permiso ya existe")
        db.refresh(obj); return _payload(obj)

    @app.patch("/api/admin/permisos/{permission_id}")
    def update_permission(permission_id: int, body: PermissionInput, user: Usuario = Depends(require_permission("roles.configurar")), db: Session = Depends(get_db)):
        obj = db.get(Permiso, permission_id)
        if not obj: raise HTTPException(status_code=404, detail="Permiso no encontrado")
        obj.clave = body.clave.strip(); obj.descripcion = body.descripcion.strip()
        try: db.commit()
        except IntegrityError: db.rollback(); raise HTTPException(status_code=409, detail="La clave de permiso ya existe")
        db.refresh(obj); return _payload(obj)

    @app.delete("/api/admin/permisos/{permission_id}", status_code=204)
    def delete_permission(permission_id: int, user: Usuario = Depends(require_permission("roles.configurar")), db: Session = Depends(get_db)):
        obj = db.get(Permiso, permission_id)
        if not obj: raise HTTPException(status_code=404, detail="Permiso no encontrado")
        if obj.roles: raise HTTPException(status_code=409, detail="Quita el permiso de todos los roles antes de eliminarlo")
        db.delete(obj); db.commit(); return Response(status_code=204)

    @app.put("/api/admin/roles/{role_id}/permisos")
    def update_role_permissions(role_id: int, body: RolePermissionsInput, user: Usuario = Depends(require_permission("roles.configurar")), db: Session = Depends(get_db)):
        role = db.query(Rol).filter_by(id=role_id).one_or_none()
        if not role:
            raise HTTPException(status_code=404, detail="Rol no encontrado")
        perms = db.query(Permiso).filter(Permiso.id.in_(body.permiso_ids)).all() if body.permiso_ids else []
        if len(perms) != len(set(body.permiso_ids)):
            raise HTTPException(status_code=422, detail="Uno o más permisos no existen")
        if role.nombre == "ADMIN":
            critical = {"roles.configurar", "usuarios.ver", "usuarios.ver_todas", "usuarios.crear", "usuarios.editar", "sucursales.ver_todas", "sucursales.crear"}
            granted = {p.clave for p in perms}
            if not critical.issubset(granted):
                raise HTTPException(status_code=409, detail="La matriz de ADMIN debe conservar los permisos de administración críticos")
        role.permisos = perms
        db.commit()
        return {"id": role.id, "nombre": role.nombre, "permisos": sorted(p.clave for p in role.permisos)}

    @app.get("/api/admin/usuarios")
    def list_users(user: Usuario = Depends(require_permission("usuarios.ver")), db: Session = Depends(get_db)):
        query = db.query(Usuario).options(joinedload(Usuario.rol), joinedload(Usuario.sucursal))
        if "usuarios.ver_todas" not in permissions_for(db, user):
            query = query.filter(Usuario.sucursal_id == user.sucursal_id)
        rows = query.order_by(Usuario.sucursal_id, Usuario.id).all()
        return [_user_payload(x, permissions_for(db, x)) for x in rows]

    @app.post("/api/admin/usuarios", status_code=201)
    def create_user(body: UserCreate, user: Usuario = Depends(require_permission("usuarios.crear")), db: Session = Depends(get_db)):
        email = body.correo.strip().lower()
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            raise HTTPException(status_code=422, detail="Correo inválido")
        branch_id = body.sucursal_id or user.sucursal_id
        if branch_id != user.sucursal_id and "sucursales.ver_todas" not in permissions_for(db, user):
            raise HTTPException(status_code=403, detail="No puede asignar usuarios a otra sucursal")
        if not db.get(Sucursal, branch_id) or not db.get(Rol, body.rol_id):
            raise HTTPException(status_code=422, detail="Sucursal o rol inválido")
        obj = Usuario(nombre=body.nombre.strip(), correo=email, password_hash=hash_password(body.password), rol_id=body.rol_id, sucursal_id=branch_id, activo=True)
        db.add(obj)
        try:
            db.commit()
        except IntegrityError:
            db.rollback(); raise HTTPException(status_code=409, detail="El correo ya está registrado")
        db.refresh(obj)
        return _user_payload(obj, permissions_for(db, obj))

    @app.patch("/api/admin/usuarios/{user_id}")
    def update_user(user_id: int, body: UserUpdate, user: Usuario = Depends(require_permission("usuarios.editar")), db: Session = Depends(get_db)):
        query = db.query(Usuario).filter(Usuario.id == user_id)
        if "usuarios.ver_todas" not in permissions_for(db, user):
            query = query.filter(Usuario.sucursal_id == user.sucursal_id)
        obj = query.one_or_none()
        if not obj:
            raise HTTPException(status_code=404, detail="Usuario no encontrado en el alcance permitido")
        changes = body.model_dump(exclude_unset=True)
        if "correo" in changes and changes["correo"]:
            changes["correo"] = changes["correo"].strip().lower()
        for key, value in changes.items():
            if key == "rol_id" and value and not db.get(Rol, value):
                raise HTTPException(status_code=422, detail="Rol inválido")
            setattr(obj, key, value)
        try:
            db.commit()
        except IntegrityError:
            db.rollback(); raise HTTPException(status_code=409, detail="Correo duplicado")
        db.refresh(obj)
        return _user_payload(obj, permissions_for(db, obj))

    # Public menu is chain-wide. Operational tables, inventory and orders are branch-scoped.
    @app.get("/api/categorias")
    def categories(user: Usuario = Depends(require_permission("productos.ver")), db: Session = Depends(get_db)):
        return [_payload(x) for x in db.query(Categoria).order_by(Categoria.nombre).all()]

    @app.post("/api/categorias", status_code=201)
    def create_category(body: CategoriaInput, user: Usuario = Depends(require_permission("productos.crear")), db: Session = Depends(get_db)):
        obj = Categoria(nombre=body.nombre.strip())
        db.add(obj)
        try:
            db.commit()
        except IntegrityError:
            db.rollback(); raise HTTPException(status_code=409, detail="La categoría ya existe")
        db.refresh(obj); return _payload(obj)

    @app.patch("/api/categorias/{category_id}")
    def update_category(category_id: int, body: CategoriaInput, user: Usuario = Depends(require_permission("productos.editar")), db: Session = Depends(get_db)):
        obj = db.get(Categoria, category_id)
        if not obj: raise HTTPException(status_code=404, detail="Categoría no encontrada")
        obj.nombre = body.nombre.strip()
        try: db.commit()
        except IntegrityError: db.rollback(); raise HTTPException(status_code=409, detail="Nombre de categoría duplicado")
        db.refresh(obj); return _payload(obj)

    @app.delete("/api/categorias/{category_id}", status_code=204)
    def delete_category(category_id: int, user: Usuario = Depends(require_permission("productos.eliminar")), db: Session = Depends(get_db)):
        obj = db.get(Categoria, category_id)
        if not obj: raise HTTPException(status_code=404, detail="Categoría no encontrada")
        if db.query(Producto).filter_by(categoria_id=obj.id).count(): raise HTTPException(status_code=409, detail="La categoría tiene productos asociados")
        db.delete(obj); db.commit(); return Response(status_code=204)

    @app.get("/api/productos")
    def products(user: Usuario = Depends(require_permission("productos.ver")), db: Session = Depends(get_db), activos: bool = True):
        q = db.query(Producto).options(joinedload(Producto.categoria))
        if activos: q = q.filter(Producto.activo.is_(True))
        return [{**_payload(x), "categoria": x.categoria.nombre if x.categoria else None} for x in q.order_by(Producto.nombre).all()]

    @app.post("/api/productos", status_code=201)
    def create_product(body: ProductoInput, user: Usuario = Depends(require_permission("productos.crear")), db: Session = Depends(get_db)):
        if not db.get(Categoria, body.categoria_id): raise HTTPException(status_code=422, detail="Categoría inválida")
        obj = Producto(**body.model_dump())
        db.add(obj); db.commit(); db.refresh(obj); return _payload(obj)

    @app.patch("/api/productos/{product_id}")
    def update_product(product_id: int, body: ProductoInput, user: Usuario = Depends(require_permission("productos.editar")), db: Session = Depends(get_db)):
        obj = db.get(Producto, product_id)
        if not obj: raise HTTPException(status_code=404, detail="Producto no encontrado")
        for key, value in body.model_dump().items(): setattr(obj, key, value)
        db.commit(); db.refresh(obj); return _payload(obj)

    @app.delete("/api/productos/{product_id}", status_code=204)
    def disable_product(product_id: int, user: Usuario = Depends(require_permission("productos.eliminar")), db: Session = Depends(get_db)):
        obj = db.get(Producto, product_id)
        if not obj: raise HTTPException(status_code=404, detail="Producto no encontrado")
        obj.activo = False; db.commit(); return Response(status_code=204)

    @app.get("/api/productos/{product_id}/receta")
    def get_recipe(product_id: int, user: Usuario = Depends(require_permission("productos.ver")), db: Session = Depends(get_db)):
        product = db.get(Producto, product_id)
        if not product: raise HTTPException(status_code=404, detail="Producto no encontrado")
        rows = db.execute(select(recetas.c.insumo_id, recetas.c.cantidad).where(recetas.c.producto_id == product_id)).all()
        result = []
        for ingredient_id, qty in rows:
            ingredient = db.get(Insumo, ingredient_id)
            result.append({"insumo_id": ingredient_id, "insumo": ingredient.nombre, "unidad_medida": ingredient.unidad_medida, "cantidad": float(qty)})
        return result

    @app.put("/api/productos/{product_id}/receta")
    def replace_recipe(product_id: int, body: RecetaReplace, user: Usuario = Depends(require_permission("recetas.configurar")), db: Session = Depends(get_db)):
        if not db.get(Producto, product_id): raise HTTPException(status_code=404, detail="Producto no encontrado")
        for line in body.insumos:
            if not db.get(Insumo, line.insumo_id): raise HTTPException(status_code=422, detail=f"Insumo {line.insumo_id} inválido")
        db.execute(recetas.delete().where(recetas.c.producto_id == product_id))
        for line in body.insumos:
            db.execute(recetas.insert().values(producto_id=product_id, insumo_id=line.insumo_id, cantidad=line.cantidad))
        db.commit()
        return get_recipe(product_id, user, db)

    @app.get("/api/insumos")
    def ingredients(user: Usuario = Depends(require_permission("insumos.ver")), db: Session = Depends(get_db)):
        return [_payload(x) for x in db.query(Insumo).order_by(Insumo.nombre).all()]

    @app.post("/api/insumos", status_code=201)
    def create_ingredient(body: InsumoInput, user: Usuario = Depends(require_permission("insumos.crear")), db: Session = Depends(get_db)):
        obj = Insumo(**body.model_dump())
        try:
            db.add(obj); db.flush()
            for branch in db.query(Sucursal).filter(Sucursal.activa.is_(True)).all():
                db.add(Inventario(sucursal_id=branch.id, insumo_id=obj.id, existencia=Decimal("0"), stock_minimo=Decimal("2")))
            db.commit()
        except IntegrityError:
            db.rollback(); raise HTTPException(status_code=409, detail="El insumo ya existe")
        db.refresh(obj); return _payload(obj)

    @app.patch("/api/insumos/{ingredient_id}")
    def update_ingredient(ingredient_id: int, body: InsumoInput, user: Usuario = Depends(require_permission("insumos.crear")), db: Session = Depends(get_db)):
        obj = db.get(Insumo, ingredient_id)
        if not obj: raise HTTPException(status_code=404, detail="Insumo no encontrado")
        obj.nombre = body.nombre.strip(); obj.unidad_medida = body.unidad_medida.strip()
        try: db.commit()
        except IntegrityError: db.rollback(); raise HTTPException(status_code=409, detail="El insumo ya existe")
        db.refresh(obj); return _payload(obj)

    @app.get("/api/inventario")
    def inventory(user: Usuario = Depends(require_permission("inventario.ver")), db: Session = Depends(get_db), sucursal_id: int | None = None, solo_bajo: bool = False):
        branch_id = scoped_branch(user, sucursal_id, allow_global=True)
        q = db.query(Inventario).options(joinedload(Inventario.insumo)).filter(Inventario.sucursal_id == branch_id)
        rows = q.order_by(Inventario.id).all()
        if solo_bajo: rows = [x for x in rows if Decimal(x.existencia) <= Decimal(x.stock_minimo)]
        return [{**_payload(x), "insumo": x.insumo.nombre, "unidad_medida": x.insumo.unidad_medida, "stock_bajo": Decimal(x.existencia) <= Decimal(x.stock_minimo)} for x in rows]

    @app.post("/api/inventario/{inventory_id}/movimientos", status_code=201)
    def adjust_inventory(inventory_id: int, body: InventarioAdjustment, user: Usuario = Depends(require_permission("inventario.ajustar")), db: Session = Depends(get_db), sucursal_id: int | None = None):
        branch_id = scoped_branch(user, sucursal_id, allow_global=True)
        item = db.query(Inventario).filter_by(id=inventory_id, sucursal_id=branch_id).with_for_update().one_or_none()
        if not item: raise HTTPException(status_code=404, detail="Insumo de inventario no encontrado en su sucursal")
        qty = Decimal(body.cantidad)
        if body.tipo == "AJUSTE_NEGATIVO":
            if Decimal(item.existencia) < qty: raise HTTPException(status_code=409, detail="La existencia no puede ser negativa")
            item.existencia = Decimal(item.existencia) - qty
        else:
            item.existencia = Decimal(item.existencia) + qty
        move = MovimientoInventario(inventario_id=item.id, tipo=body.tipo, cantidad=qty, usuario_id=user.id, nota=body.nota)
        db.add(move); db.commit(); db.refresh(item); db.refresh(move)
        return {"movimiento": _payload(move), "inventario": _payload(item)}

    @app.get("/api/mesas")
    def list_tables(user: Usuario = Depends(require_permission("mesas.ver")), db: Session = Depends(get_db)):
        return [_payload(x) for x in db.query(Mesa).filter(Mesa.sucursal_id == user.sucursal_id).order_by(Mesa.numero).all()]

    @app.post("/api/mesas", status_code=201)
    def create_table(body: MesaInput, user: Usuario = Depends(require_permission("mesas.configurar")), db: Session = Depends(get_db)):
        if body.estado not in {"LIBRE", "OCUPADA", "RESERVADA"}: raise HTTPException(status_code=422, detail="Estado de mesa inválido")
        obj = Mesa(sucursal_id=user.sucursal_id, **body.model_dump()); db.add(obj)
        try: db.commit()
        except IntegrityError: db.rollback(); raise HTTPException(status_code=409, detail="Ya existe ese número de mesa en la sucursal")
        db.refresh(obj); return _payload(obj)

    @app.patch("/api/mesas/{table_id}")
    def update_table(table_id: int, body: MesaInput, user: Usuario = Depends(require_permission("mesas.configurar")), db: Session = Depends(get_db)):
        obj = db.query(Mesa).filter_by(id=table_id, sucursal_id=user.sucursal_id).one_or_none()
        if not obj: raise HTTPException(status_code=404, detail="Mesa no encontrada en su sucursal")
        obj.numero, obj.capacidad = body.numero, body.capacidad
        if body.estado not in {"LIBRE", "OCUPADA", "RESERVADA"}: raise HTTPException(status_code=422, detail="Estado de mesa inválido")
        obj.estado = body.estado; db.commit(); db.refresh(obj); return _payload(obj)

    @app.delete("/api/mesas/{table_id}", status_code=204)
    def delete_table(table_id: int, user: Usuario = Depends(require_permission("mesas.configurar")), db: Session = Depends(get_db)):
        table = db.query(Mesa).filter_by(id=table_id, sucursal_id=user.sucursal_id).one_or_none()
        if not table: raise HTTPException(status_code=404, detail="Mesa no encontrada en su sucursal")
        if db.query(Orden).filter_by(mesa_id=table.id).filter(Orden.estado.notin_(["PAGADA", "CANCELADA"])).count():
            raise HTTPException(status_code=409, detail="No se puede borrar una mesa con órdenes activas")
        db.delete(table)
        try: db.commit()
        except IntegrityError: db.rollback(); raise HTTPException(status_code=409, detail="La mesa tiene historial de órdenes; puede cambiar su estado en lugar de borrarla")
        return Response(status_code=204)

    def order_view(db: Session, order: Orden) -> dict[str, Any]:
        details = db.query(DetalleOrden).options(joinedload(DetalleOrden.producto)).filter(DetalleOrden.orden_id == order.id).all()
        mesa = db.get(Mesa, order.mesa_id)
        paid = db.query(func.coalesce(func.sum(Pago.monto), 0)).filter(Pago.orden_id == order.id).scalar()
        return {
            "id": order.id, "sucursal_id": order.sucursal_id, "mesa_id": order.mesa_id,
            "mesa_numero": mesa.numero if mesa else None, "mesero_id": order.mesero_id,
            "estado": order.estado, "total": float(order.total), "pagado": float(paid or 0),
            "pendiente": float(max(Decimal("0"), Decimal(order.total) - Decimal(paid or 0))),
            "fecha": order.fecha.isoformat(),
            "detalles": [{"id": d.id, "producto_id": d.producto_id, "producto": d.producto.nombre, "cantidad": d.cantidad, "precio_unitario": float(d.precio_unitario), "subtotal": float(Decimal(d.precio_unitario) * d.cantidad), "notas": d.notas, "estado": d.estado} for d in details],
        }

    def recalculate_order(db: Session, order: Orden):
        details = db.query(DetalleOrden).filter(DetalleOrden.orden_id == order.id, DetalleOrden.estado != "CANCELADO").all()
        order.total = money(sum((Decimal(d.precio_unitario) * d.cantidad for d in details), Decimal("0")))

    @app.get("/api/ordenes")
    def list_orders(user: Usuario = Depends(require_permission("ordenes.ver")), db: Session = Depends(get_db), estado: str | None = None):
        q = _branch_orders(db, user).order_by(Orden.fecha.desc())
        if estado: q = q.filter(Orden.estado == estado)
        return [order_view(db, x) for x in q.limit(300).all()]

    @app.post("/api/ordenes", status_code=201)
    async def create_order(body: OrdenCreate, request: Request, user: Usuario = Depends(require_permission("ordenes.crear")), db: Session = Depends(get_db)):
        table = db.query(Mesa).filter_by(id=body.mesa_id, sucursal_id=user.sucursal_id).with_for_update().one_or_none()
        if not table: raise HTTPException(status_code=404, detail="Mesa no encontrada en su sucursal")
        if table.estado != "LIBRE": raise HTTPException(status_code=409, detail="La mesa no está libre")
        order = Orden(sucursal_id=user.sucursal_id, mesa_id=table.id, mesero_id=user.id, estado="ABIERTA", total=Decimal("0"))
        db.add(order); db.flush()
        for line in body.detalles:
            product = db.query(Producto).filter_by(id=line.producto_id, activo=True).one_or_none()
            if not product: db.rollback(); raise HTTPException(status_code=404, detail=f"Producto {line.producto_id} no disponible")
            db.add(DetalleOrden(orden_id=order.id, producto_id=product.id, cantidad=line.cantidad, precio_unitario=product.precio, notas=line.notas, estado="PENDIENTE"))
        table.estado = "OCUPADA"
        db.flush(); recalculate_order(db, order); db.commit(); db.refresh(order)
        request.state.audit_ref = order.id
        result = order_view(db, order)
        await SOCKET.emit("orden_actualizada", result, room=f"sucursal:{user.sucursal_id}")
        return result

    @app.post("/api/ordenes/{order_id}/detalles", status_code=201)
    async def add_order_line(order_id: int, body: OrdenLineAdd, user: Usuario = Depends(require_permission("ordenes.editar")), db: Session = Depends(get_db)):
        order = db.query(Orden).filter_by(id=order_id, sucursal_id=user.sucursal_id).with_for_update().one_or_none()
        if not order: raise HTTPException(status_code=404, detail="Orden no encontrada en su sucursal")
        if order.estado != "ABIERTA": raise HTTPException(status_code=409, detail="Solo se editan órdenes ABIERTAS")
        product = db.query(Producto).filter_by(id=body.producto_id, activo=True).one_or_none()
        if not product: raise HTTPException(status_code=404, detail="Producto no disponible")
        db.add(DetalleOrden(orden_id=order.id, producto_id=product.id, cantidad=body.cantidad, precio_unitario=product.precio, notas=body.notas))
        db.flush(); recalculate_order(db, order); db.commit(); db.refresh(order)
        result = order_view(db, order); await SOCKET.emit("orden_actualizada", result, room=f"sucursal:{user.sucursal_id}"); return result

    @app.patch("/api/ordenes/{order_id}/detalles/{detail_id}")
    async def update_order_line(order_id: int, detail_id: int, body: OrdenLineUpdate, user: Usuario = Depends(require_permission("ordenes.editar")), db: Session = Depends(get_db)):
        order = db.query(Orden).filter_by(id=order_id, sucursal_id=user.sucursal_id).with_for_update().one_or_none()
        if not order: raise HTTPException(status_code=404, detail="Orden no encontrada en su sucursal")
        if order.estado != "ABIERTA": raise HTTPException(status_code=409, detail="Solo se modifican detalles de órdenes ABIERTAS")
        detail = db.query(DetalleOrden).filter_by(id=detail_id, orden_id=order.id).with_for_update().one_or_none()
        if not detail: raise HTTPException(status_code=404, detail="Platillo no encontrado")
        if detail.estado == "CANCELADO": raise HTTPException(status_code=409, detail="El platillo está cancelado")
        changes = body.model_dump(exclude_unset=True)
        for key, value in changes.items(): setattr(detail, key, value)
        recalculate_order(db, order); db.commit(); db.refresh(order)
        result = order_view(db, order); await SOCKET.emit("orden_actualizada", result, room=f"sucursal:{user.sucursal_id}"); return result

    def return_ingredients(db: Session, detail: DetalleOrden, order: Orden, user_id: int, reason: str):
        if detail.stock_reintegrado:
            return
        recipe_rows = db.execute(select(recetas.c.insumo_id, recetas.c.cantidad).where(recetas.c.producto_id == detail.producto_id)).all()
        for ingredient_id, recipe_qty in recipe_rows:
            inv = db.query(Inventario).filter_by(sucursal_id=order.sucursal_id, insumo_id=ingredient_id).with_for_update().one_or_none()
            if not inv:
                raise HTTPException(status_code=409, detail=f"No hay registro de inventario para insumo {ingredient_id}")
            qty = Decimal(recipe_qty) * detail.cantidad
            inv.existencia = Decimal(inv.existencia) + qty
            db.add(MovimientoInventario(inventario_id=inv.id, tipo="REINTEGRO", cantidad=qty, orden_id=order.id, usuario_id=user_id, nota=reason))
        detail.stock_reintegrado = True

    @app.delete("/api/ordenes/{order_id}/detalles/{detail_id}")
    async def cancel_order_line(order_id: int, detail_id: int, user: Usuario = Depends(require_permission("ordenes.cancelar")), db: Session = Depends(get_db)):
        order = db.query(Orden).filter_by(id=order_id, sucursal_id=user.sucursal_id).with_for_update().one_or_none()
        if not order: raise HTTPException(status_code=404, detail="Orden no encontrada")
        if order.estado not in {"ABIERTA", "EN_COCINA"}: raise HTTPException(status_code=409, detail="Ya no puede cancelarse el platillo en este estado")
        detail = db.query(DetalleOrden).filter_by(id=detail_id, orden_id=order.id).with_for_update().one_or_none()
        if not detail: raise HTTPException(status_code=404, detail="Platillo no encontrado")
        if detail.estado == "CANCELADO": raise HTTPException(status_code=409, detail="El platillo ya está cancelado")
        if order.estado == "EN_COCINA": return_ingredients(db, detail, order, user.id, "Reintegro por cancelación de platillo")
        detail.estado = "CANCELADO"; recalculate_order(db, order); db.commit(); db.refresh(order)
        result = order_view(db, order); await SOCKET.emit("platillo_actualizado", {"orden": result, "detalle_id": detail.id, "estado": "CANCELADO"}, room=f"sucursal:{user.sucursal_id}"); return result

    @app.post("/api/ordenes/{order_id}/enviar-cocina")
    async def send_order_to_kitchen(order_id: int, user: Usuario = Depends(require_permission("ordenes.enviar_cocina")), db: Session = Depends(get_db)):
        order = db.query(Orden).filter_by(id=order_id, sucursal_id=user.sucursal_id).with_for_update().one_or_none()
        if not order: raise HTTPException(status_code=404, detail="Orden no encontrada")
        if order.estado != "ABIERTA": raise HTTPException(status_code=409, detail="La orden no está ABIERTA")
        details = db.query(DetalleOrden).filter(DetalleOrden.orden_id == order.id, DetalleOrden.estado != "CANCELADO").all()
        if not details: raise HTTPException(status_code=409, detail="La orden no tiene platillos activos")
        # Aggregate every recipe line by ingredient before validating stock. This prevents
        # duplicate ingredients across multiple dishes from bypassing the stock check.
        requirements: dict[int, Decimal] = {}
        for detail in details:
            recipe = db.execute(select(recetas.c.insumo_id, recetas.c.cantidad).where(recetas.c.producto_id == detail.producto_id)).all()
            if not recipe:
                raise HTTPException(status_code=409, detail=f"El producto {detail.producto_id} no tiene receta configurada")
            for ingredient_id, recipe_qty in recipe:
                requirements[ingredient_id] = requirements.get(ingredient_id, Decimal("0")) + Decimal(recipe_qty) * detail.cantidad
        deductions: list[tuple[Inventario, Decimal]] = []
        for ingredient_id, amount in requirements.items():
            inv = db.query(Inventario).filter_by(sucursal_id=order.sucursal_id, insumo_id=ingredient_id).with_for_update().one_or_none()
            if not inv:
                raise HTTPException(status_code=409, detail=f"Falta inventario para insumo {ingredient_id}")
            if Decimal(inv.existencia) < amount:
                ingredient = db.get(Insumo, ingredient_id)
                raise HTTPException(status_code=409, detail=f"Stock insuficiente de {ingredient.nombre}: disponible {inv.existencia}, requerido {amount}")
            deductions.append((inv, amount))
        # Commit all deductions and movements together; a failure rolls back the whole order.
        for inv, amount in deductions:
            inv.existencia = Decimal(inv.existencia) - amount
            db.add(MovimientoInventario(inventario_id=inv.id, tipo="SALIDA", cantidad=amount, orden_id=order.id, usuario_id=user.id, nota="Consumo agregado de recetas al enviar a cocina"))
        order.estado = "EN_COCINA"
        db.commit(); db.refresh(order)
        result = order_view(db, order)
        await SOCKET.emit("orden_nueva", result, room=f"sucursal:{user.sucursal_id}")
        await SOCKET.emit("orden_actualizada", result, room=f"sucursal:{user.sucursal_id}")
        return result

    @app.post("/api/ordenes/{order_id}/servir")
    async def serve_order(order_id: int, user: Usuario = Depends(require_permission("ordenes.editar")), db: Session = Depends(get_db)):
        order = db.query(Orden).filter_by(id=order_id, sucursal_id=user.sucursal_id).with_for_update().one_or_none()
        if not order: raise HTTPException(status_code=404, detail="Orden no encontrada")
        if order.estado != "EN_COCINA": raise HTTPException(status_code=409, detail="Solo se sirven órdenes en cocina")
        active = db.query(DetalleOrden).filter(DetalleOrden.orden_id == order.id, DetalleOrden.estado != "CANCELADO").all()
        if not active or any(x.estado != "LISTO" for x in active): raise HTTPException(status_code=409, detail="Todos los platillos deben estar LISTOS")
        order.estado = "SERVIDA"; db.commit(); db.refresh(order)
        result = order_view(db, order); await SOCKET.emit("orden_actualizada", result, room=f"sucursal:{user.sucursal_id}"); return result

    @app.post("/api/ordenes/{order_id}/cancelar")
    async def cancel_order(order_id: int, body: OrderCancelInput, user: Usuario = Depends(require_permission("ordenes.cancelar")), db: Session = Depends(get_db)):
        order = db.query(Orden).filter_by(id=order_id, sucursal_id=user.sucursal_id).with_for_update().one_or_none()
        if not order: raise HTTPException(status_code=404, detail="Orden no encontrada")
        if order.estado not in {"ABIERTA", "EN_COCINA"}: raise HTTPException(status_code=409, detail="La orden no puede cancelarse en su estado actual")
        if order.estado == "EN_COCINA":
            details = db.query(DetalleOrden).filter(DetalleOrden.orden_id == order.id, DetalleOrden.estado != "CANCELADO").with_for_update().all()
            for detail in details:
                return_ingredients(db, detail, order, user.id, body.motivo)
                detail.estado = "CANCELADO"
        order.estado = "CANCELADA"
        table = db.query(Mesa).filter_by(id=order.mesa_id, sucursal_id=user.sucursal_id).with_for_update().one()
        table.estado = "LIBRE"
        db.commit(); db.refresh(order)
        result = order_view(db, order); await SOCKET.emit("orden_actualizada", result, room=f"sucursal:{user.sucursal_id}"); return result

    @app.get("/api/ordenes/{order_id}")
    def get_order(order_id: int, user: Usuario = Depends(require_permission("ordenes.ver")), db: Session = Depends(get_db)):
        order = db.query(Orden).filter_by(id=order_id, sucursal_id=user.sucursal_id).one_or_none()
        if not order: raise HTTPException(status_code=404, detail="Orden no encontrada en su sucursal")
        return order_view(db, order)

    @app.get("/api/inventario/movimientos")
    def inventory_movements(user: Usuario = Depends(require_permission("inventario.ver")), db: Session = Depends(get_db), limite: int = Query(default=100, ge=1, le=500)):
        ids = [row.id for row in db.query(Inventario.id).filter_by(sucursal_id=user.sucursal_id).all()]
        rows = db.query(MovimientoInventario).options(joinedload(MovimientoInventario.inventario)).filter(MovimientoInventario.inventario_id.in_(ids)).order_by(MovimientoInventario.fecha.desc()).limit(limite).all()
        return [{**_payload(x), "insumo_id": x.inventario.insumo_id, "sucursal_id": x.inventario.sucursal_id} for x in rows]

    @app.get("/api/cocina/ordenes")
    def kitchen_orders(user: Usuario = Depends(require_permission("cocina.ver")), db: Session = Depends(get_db)):
        orders = db.query(Orden).options(joinedload(Orden.mesa)).filter(Orden.sucursal_id == user.sucursal_id, Orden.estado == "EN_COCINA").order_by(Orden.fecha.asc()).all()
        return [order_view(db, x) for x in orders]

    @app.patch("/api/cocina/detalles/{detail_id}/estado")
    async def update_dish_state(detail_id: int, body: DishStateInput, user: Usuario = Depends(require_permission("cocina.actualizar")), db: Session = Depends(get_db)):
        detail = db.query(DetalleOrden).options(joinedload(DetalleOrden.orden)).filter(DetalleOrden.id == detail_id).with_for_update(of=DetalleOrden).one_or_none()
        if not detail or detail.orden.sucursal_id != user.sucursal_id: raise HTTPException(status_code=404, detail="Platillo no encontrado en su sucursal")
        if detail.orden.estado != "EN_COCINA": raise HTTPException(status_code=409, detail="La orden no está activa en cocina")
        allowed = {"PENDIENTE": {"PREPARANDO"}, "PREPARANDO": {"LISTO"}}
        if body.estado not in allowed.get(detail.estado, set()): raise HTTPException(status_code=409, detail=f"Transición no permitida: {detail.estado} -> {body.estado}")
        detail.estado = body.estado; db.commit(); db.refresh(detail)
        result = {"detalle_id": detail.id, "orden_id": detail.orden_id, "producto_id": detail.producto_id, "estado": detail.estado, "sucursal_id": detail.orden.sucursal_id}
        await SOCKET.emit("platillo_actualizado", result, room=f"sucursal:{user.sucursal_id}")
        return result

    @app.get("/api/pagos/ordenes-abiertas")
    def payment_orders(user: Usuario = Depends(require_permission("pagos.ver")), db: Session = Depends(get_db)):
        orders = _branch_orders(db, user).filter(Orden.estado == "SERVIDA").order_by(Orden.fecha.asc()).all()
        return [order_view(db, x) for x in orders]

    @app.post("/api/ordenes/{order_id}/pagos", status_code=201)
    async def create_payment(order_id: int, body: PaymentInput, user: Usuario = Depends(require_permission("pagos.crear")), db: Session = Depends(get_db)):
        order = db.query(Orden).filter_by(id=order_id, sucursal_id=user.sucursal_id).with_for_update().one_or_none()
        if not order: raise HTTPException(status_code=404, detail="Orden no encontrada en su sucursal")
        if order.estado != "SERVIDA": raise HTTPException(status_code=409, detail="Solo puede cobrarse una orden SERVIDA")
        current_paid = Decimal(db.query(func.coalesce(func.sum(Pago.monto), 0)).filter(Pago.orden_id == order.id).scalar() or 0)
        amount = money(body.monto)
        total = money(order.total)
        if current_paid + amount > total:
            raise HTTPException(status_code=409, detail=f"El pago excede el saldo pendiente de {money(total-current_paid)}")
        payment = Pago(orden_id=order.id, cajero_id=user.id, metodo=body.metodo, monto=amount)
        db.add(payment); db.flush()
        new_paid = current_paid + amount
        closed = new_paid == total
        if closed:
            order.estado = "PAGADA"
            table = db.query(Mesa).filter_by(id=order.mesa_id, sucursal_id=user.sucursal_id).with_for_update().one()
            table.estado = "LIBRE"
        db.commit(); db.refresh(payment); db.refresh(order)
        result = {"pago": _payload(payment), "orden": order_view(db, order), "total_pagado": float(new_paid), "saldo": float(total-new_paid), "cerrada": closed}
        await SOCKET.emit("orden_actualizada", result["orden"], room=f"sucursal:{user.sucursal_id}")
        return result

    @app.get("/api/reportes/ventas")
    def sales_report(user: Usuario = Depends(require_permission("reportes.ver")), db: Session = Depends(get_db), desde: datetime | None = None, hasta: datetime | None = None, sucursal_id: int | None = None):
        perms = set(permissions_for(db, user))
        if sucursal_id is not None and sucursal_id != user.sucursal_id and "reportes.ver_todas" not in perms:
            raise HTTPException(status_code=403, detail="No tiene permiso para consultar otra sucursal")
        q = db.query(Sucursal.id, Sucursal.nombre, func.coalesce(func.sum(Pago.monto), 0).label("ventas"), func.count(Pago.id).label("pagos")).outerjoin(Orden, and_(Orden.sucursal_id == Sucursal.id, Orden.estado == "PAGADA")).outerjoin(Pago, Pago.orden_id == Orden.id)
        if desde: q = q.filter((Pago.fecha >= desde) | (Pago.id.is_(None)))
        if hasta: q = q.filter((Pago.fecha <= hasta) | (Pago.id.is_(None)))
        if "reportes.ver_todas" not in perms:
            q = q.filter(Sucursal.id == user.sucursal_id)
        elif sucursal_id:
            q = q.filter(Sucursal.id == sucursal_id)
        rows = q.group_by(Sucursal.id, Sucursal.nombre).order_by(Sucursal.id).all()
        return [{"sucursal_id": row.id, "sucursal": row.nombre, "ventas": float(row.ventas or 0), "pagos": int(row.pagos or 0)} for row in rows]

    @app.get("/api/reportes/productos-mas-vendidos")
    def top_products(user: Usuario = Depends(require_permission("reportes.ver")), db: Session = Depends(get_db), limite: int = Query(default=10, ge=1, le=100)):
        q = db.query(Producto.id, Producto.nombre, func.sum(DetalleOrden.cantidad).label("unidades"), func.sum(DetalleOrden.cantidad * DetalleOrden.precio_unitario).label("importe")).join(DetalleOrden, DetalleOrden.producto_id == Producto.id).join(Orden, Orden.id == DetalleOrden.orden_id).filter(Orden.sucursal_id == user.sucursal_id, Orden.estado == "PAGADA", DetalleOrden.estado != "CANCELADO").group_by(Producto.id, Producto.nombre).order_by(func.sum(DetalleOrden.cantidad).desc()).limit(limite).all()
        return [{"producto_id": x.id, "producto": x.nombre, "unidades": int(x.unidades or 0), "importe": float(x.importe or 0)} for x in q]

    @app.get("/api/reportes/stock-bajo")
    def low_stock_report(user: Usuario = Depends(require_permission("reportes.ver")), db: Session = Depends(get_db)):
        rows = db.query(Inventario).options(joinedload(Inventario.insumo)).filter(Inventario.sucursal_id == user.sucursal_id, Inventario.existencia <= Inventario.stock_minimo).order_by(Inventario.existencia.asc()).all()
        return [{"inventario_id": x.id, "insumo_id": x.insumo_id, "insumo": x.insumo.nombre, "existencia": float(x.existencia), "stock_minimo": float(x.stock_minimo), "unidad": x.insumo.unidad_medida} for x in rows]

    @app.get("/api/reportes/ventas.csv")
    def sales_csv(user: Usuario = Depends(require_permission("reportes.ver")), db: Session = Depends(get_db), desde: datetime | None = None, hasta: datetime | None = None, sucursal_id: int | None = None):
        rows = sales_report(user, db, desde, hasta, sucursal_id)
        output = io.StringIO(); writer = csv.DictWriter(output, fieldnames=["sucursal_id", "sucursal", "ventas", "pagos"])
        writer.writeheader(); writer.writerows(rows)
        return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=comandapro_ventas.csv"})

    @app.get("/api/reportes/ventas.xlsx")
    def sales_xlsx(user: Usuario = Depends(require_permission("reportes.ver")), db: Session = Depends(get_db), desde: datetime | None = None, hasta: datetime | None = None, sucursal_id: int | None = None):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        rows = sales_report(user, db, desde, hasta, sucursal_id)
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Ventas por sucursal"
        sheet.append(["ID sucursal", "Sucursal", "Ventas (MXN)", "Número de pagos"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D503B")
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append([row["sucursal_id"], row["sucursal"], row["ventas"], row["pagos"]])
        sheet.column_dimensions["A"].width = 15; sheet.column_dimensions["B"].width = 30
        sheet.column_dimensions["C"].width = 18; sheet.column_dimensions["D"].width = 20
        for cells in sheet.iter_rows(min_row=2, min_col=3, max_col=3): cells[0].number_format = '"$"#,##0.00'
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        output = io.BytesIO(); workbook.save(output); output.seek(0)
        return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=comandapro_ventas.xlsx"})

    @app.get("/api/bitacora")
    def audit_log(user: Usuario = Depends(require_permission("bitacora.ver")), db: Session = Depends(get_db), limite: int = Query(default=100, ge=1, le=500), accion: str | None = None, entidad: str | None = None, usuario_id: int | None = None, desde: datetime | None = None, hasta: datetime | None = None):
        user_ids = [x.id for x in db.query(Usuario.id).filter(Usuario.sucursal_id == user.sucursal_id).all()]
        q = db.query(Bitacora).options(joinedload(Bitacora.usuario)).filter(Bitacora.usuario_id.in_(user_ids))
        if accion: q = q.filter(Bitacora.accion == accion.upper())
        if entidad: q = q.filter(Bitacora.entidad == entidad)
        if usuario_id: q = q.filter(Bitacora.usuario_id == usuario_id)
        if desde: q = q.filter(Bitacora.fecha >= desde)
        if hasta: q = q.filter(Bitacora.fecha <= hasta)
        rows = q.order_by(Bitacora.fecha.desc()).limit(limite).all()
        return [{**_payload(x), "usuario": x.usuario.nombre if x.usuario else None} for x in rows]

    return app


api = create_app()

# Render: sirve el frontend compilado en el mismo origen para mantener
# /api y Socket.IO sin depender de un proxy o de otro servicio gratuito.
if os.getenv("COMANDAPRO_SERVE_FRONTEND", "").lower() == "true":
    from fastapi.staticfiles import StaticFiles
    api.mount("/", StaticFiles(directory="/app/frontend-dist", html=True), name="frontend")


@SOCKET.event
async def connect(sid, environ, auth):
    global _SOCKET_SESSION_FACTORY
    try:
        token = (auth or {}).get("token")
        claims = decode_access_token(token or "")
        if _SOCKET_SESSION_FACTORY is None:
            return False
        with _SOCKET_SESSION_FACTORY() as db:
            user = db.query(Usuario).filter_by(id=int(claims["sub"]), activo=True).one_or_none()
            if not user:
                return False
            await SOCKET.save_session(sid, {"user_id": user.id, "sucursal_id": user.sucursal_id})
            await SOCKET.enter_room(sid, f"sucursal:{user.sucursal_id}")
        return True
    except Exception:
        return False


@SOCKET.event
async def disconnect(sid):
    return None


@api.on_event("startup")
def configure_socket_session_factory():
    global _SOCKET_SESSION_FACTORY
    _SOCKET_SESSION_FACTORY = api.state.SessionLocal


app = socketio.ASGIApp(SOCKET, other_asgi_app=api, socketio_path="socket.io")
